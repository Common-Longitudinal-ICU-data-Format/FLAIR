#!/usr/bin/env python3
"""Build docs/paper/main_v3.docx from main_v3.tex.

This is a v3-specific builder. It does not read or write the v2 manuscript files.
The parser intentionally handles the simple LaTeX subset used in main_v3.tex
and renders the manuscript tables natively with python-docx.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import flair_paper_data as D

HERE = Path(__file__).resolve().parent
SRC = HERE / "main_v3.tex"
FIG = HERE / "figures"
OUT = HERE / "main_v3.docx"
BIB = HERE / "references.bib"

TITLE = "FLAIR: A Federated ICU Benchmark for External Validation on the Common Longitudinal ICU Format"
AUTHORS = "CLIF Consortium · FLAIR Working Group"
TITLE_PAGE_TODOS = [
    "Author affiliations: [AUTHOR: complete at submission]",
    "Corresponding author: [AUTHOR: complete at submission]",
    "Manuscript type: Datasets, Benchmarks, and Protocols",
]

CITE_NUMBERS: dict[str, str] = {}
REFERENCES: list[str] = []

def _consortium_rows() -> list[list[str]]:
    rows = []
    for s in D.SITE_ORDER:
        i = D.ALL_SITES[s]
        if i.get("observed"):
            rows.append([i["label"], f"{i['stays']:,}", f"{i['female_pct']:.1f}",
                         str(i["age_median"]), str(i["sofa_median"])])
        else:
            rows.append([i["label"], "---", "---", "---", "---"])
    return rows


def _leaderboard_rows() -> list[list[str]]:
    rows = []
    for s in D.SITE_ORDER:
        i = D.ALL_SITES[s]
        if i.get("observed"):
            rows.append([i["label"]] + [D.fmt_ci(s, t) for t in D.TASK_IDS])
        else:
            rows.append([i["label"]] + ["---"] * len(D.TASK_IDS))
    rows.append(["External median"] + [f"{D.external_median(t):.3f}" for t in D.TASK_IDS])
    return rows


_TASK_LONG = {
    "T1": "T1, ICU daily mortality", "T2": "T2, LTACH discharge",
    "T3": "T3, extubation failure 24 h", "T4": "T4, sepsis 6 h",
    "T5": "T5, unplanned ICU readmission",
}


def _mimic_full_rows() -> list[list[str]]:
    rows = []
    for tid in D.TASK_IDS:
        fm = D.FULL_METRICS["mimic"][tid]
        cal = fm["calibration"]
        slope = f"{cal[0]:.2f}" if cal else "---"
        intc = f"{cal[1]:.2f}" if cal else "---"
        brier = f"{cal[2]:.3f}" if cal else "---"
        rows.append([
            _TASK_LONG[tid], fm["mode"], f"{fm['n']:,}", f"{fm['prevalence']:.3f}",
            D.fmt_ci("mimic", tid), f"{fm['auprc']:.3f}", slope, intc, brier,
        ])
    return rows


TABLES = {
    # ---- main text ----
    "tab:taskspec": {
        "caption": (
            "Table 1. The five FLAIR benchmark tasks. Prevalence is the MIMIC-CLIF "
            "positive rate on the scored test set (first landmark for T1-T2; per stay "
            "for T3-T5). Clinical thresholds and directions are single-sourced in the "
            "package policy registry and folded into each task's audit metadata. "
            "Source: FLAIR task code and baseline report artifacts."
        ),
        "header": ["Task", "Outcome", "Prediction grid", "Report mode", "Threshold", "Prev."],
        "rows": [
            [t[0], t[1], t[5], t[2], f"{t[3]:.2f} ({t[4]})",
             f"{D.FULL_METRICS['mimic'][t[0]]['prevalence']:.3f}"]
            for t in D.TASKS
        ],
    },
    "tab:consortium": {
        "caption": (
            "Table 2. Federated consortium cohort characteristics. Age is the "
            "median in years; SOFA is the median hospital-maximum score. Four sites "
            "have reported observed federated runs; four additional consortium sites "
            "are onboarding (shown as placeholder rows) and have not yet returned "
            "results. Per-site cell suppression at n<10 is applied. Source: FLAIR "
            "baseline report artifacts, single-sourced through flair_paper_data.py."
        ),
        "header": ["Site", "ICU stays", "Female %", "Age", "SOFA"],
        "rows": _consortium_rows(),
    },
    # ---- supplement ----
    "tab:supp-baseline-full": {
        "caption": (
            "Table S1. Full observed MIMIC-CLIF internal baseline metrics for the "
            "XGBoost count-feature model. Landmark tasks (T1, T2) report the pooled "
            "dynamic AUC and pooled calibration across landmarks. Calibration is "
            "omitted for T4 by design under peak scoring. Source: FLAIR baseline "
            "report artifacts."
        ),
        "header": ["Task", "Mode", "n", "Prev.", "AUROC (95% CI)", "AUPRC", "Cal. slope", "Cal. int.", "Brier"],
        "rows": _mimic_full_rows(),
    },
    "tab:supp-leaderboard": {
        "caption": (
            "Table S2. Per-site AUROC (95% CI) for the frozen baseline, mode-matched "
            "per task; landmark tasks (T1, T2) report the pooled dynamic AUC. Four "
            "sites are observed; four consortium sites are onboarding (placeholder "
            "rows). External median is over the three observed external sites. Source: "
            "FLAIR baseline report artifacts, single-sourced through flair_paper_data.py."
        ),
        "header": ["Site", "T1 (95% CI)", "T2 (95% CI)", "T3 (95% CI)", "T4 (95% CI)", "T5 (95% CI)"],
        "rows": _leaderboard_rows(),
    },
    "tab:mimic-snapshot": {
        "caption": (
            "Table S3. Observed MIMIC-CLIF per-task cohort snapshot. Counts are "
            "eligible stitched encounters (hospitalization_join_id) per task; T1 and "
            "T2 share the daily-landmark cohort, while T3-T5 apply task-specific "
            "eligibility. Charlson scores are missing for roughly half of stays in the "
            "MIMIC-CLIF export. Source: FLAIR baseline report artifacts (per-task table1.json)."
        ),
        "header": ["Measure", "T1 mortality", "T2 LTACH", "T3 extubation", "T4 sepsis", "T5 readmission"],
        "rows": [
            ["Eligible hospitalizations", "72,778", "72,778", "26,770", "79,772", "67,221"],
            ["Sex, F/M", "31,409/41,369", "31,409/41,369", "10,032/16,738", "35,415/44,357", "29,195/38,026"],
            ["Age, median (IQR)", "67 (55-77)", "67 (55-77)", "66 (56-75)", "67 (55-77)", "66 (55-77)"],
            ["SOFA hospital maximum, median (IQR)", "5 (2-8)", "5 (2-8)", "7 (5-10)", "4 (2-7)", "4 (2-7)"],
            ["Charlson score, median (IQR)", "2 (0-4)", "2 (0-4)", "2 (0-4)", "2 (0-4)", "2 (0-4)"],
        ],
    },
    "tab:comparators": {
        "caption": (
            "Table S4. FLAIR compared with adjacent federated or model-to-data "
            "infrastructures. Intended as orientation, not a formal ranking. Source: authors."
        ),
        "header": ["Feature", "OHDSI/LEGEND", "MedPerf", "FeTS", "OpenSAFELY", "FLAIR"],
        "rows": [
            ["Data standard", "OMOP", "task specific", "imaging", "TPP/EMIS", "CLIF"],
            ["Federation unit", "SQL/coefs", "container", "gradients", "analytic code", "model/code execution"],
            ["Main use", "network studies", "imaging benchmarks", "federated training", "secure analysis", "ICU EHR benchmarks"],
            ["Hosted-model portal", "no", "yes", "yes", "no", "yes"],
            ["Patient rows shared", "no", "no", "no", "no", "no"],
            ["Extensible tasks", "limited", "partial", "no", "not applicable", "yes"],
        ],
    },
}

FIGURES = {
    "fig:workflow": {
        "image": "figure1_flair_workflow_v3.png",
        "width": 6.3,
        "caption": (
            "Figure 1. The FLAIR model-to-data loop. A Model Owner develops against "
            "public MIMIC-CLIF and submits a frozen model, code, or container. The "
            "Benchmark Committee reviews it for reproducibility, clinical face "
            "validity, and security, and stamps a task audit hash. The FLAIR Portal "
            "distributes byte-identical code to participating CLIF sites, where a "
            "Site Lead or Data Steward runs it locally against site CLIF tables. "
            "Patient-level cohorts, features, and predictions remain inside the "
            "institutional boundary; only aggregate, cell-suppressed JSON report "
            "bundles return. The leaderboard foregrounds per-site variation rather "
            "than a single pooled scalar. Source: authors."
        ),
    },
    "fig:degradation": {
        "image": "figure3_degradation.png",
        "width": 6.0,
        "caption": (
            "Figure 2. Observed cross-site degradation of the frozen baseline from "
            "the MIMIC-CLIF internal holdout to the three external sites (Rush, "
            "Northwestern Medicine, University of Chicago Medicine). Left: for each "
            "task a line runs from the internal AUROC to each external site; sepsis "
            "and readmission cross below the chance band. Right: the median internal-"
            "to-external AUROC drop by task, ordered by magnitude. All four sites are "
            "observed. Landmark tasks use the pooled dynamic AUC. Source: FLAIR "
            "baseline report artifacts."
        ),
    },
    "fig:forest": {
        "image": "figure2_forest_auroc.png",
        "width": 6.5,
        "caption": (
            "Figure 3. Per-site discrimination of the frozen baseline across the CLIF "
            "consortium, one panel per task. Each marker is a site's AUROC with a 95% "
            "confidence interval; the dashed line is the MIMIC-CLIF internal reference "
            "and the shaded band marks worse-than-chance discrimination. The four "
            "observed sites (MIMIC-CLIF, Rush, Northwestern Medicine, University of "
            "Chicago Medicine) are shown in colour; four onboarding consortium sites "
            "appear as greyed placeholder rows pending federated runs. Landmark tasks "
            "(T1, T2) use the pooled dynamic AUC. Source: FLAIR baseline report "
            "artifacts, single-sourced through flair_paper_data.py."
        ),
    },
}

SUBS = [
    (r"\\todo\{([^}]*)\}", r"TODO: \1"),
    (r"\\confirm\{([^}]*)\}", r"[AUTHOR: \1]"),
    (r"\\textbf\{([^}]*)\}", r"\1"),
    (r"\\textit\{([^}]*)\}", r"\1"),
    (r"\\emph\{([^}]*)\}", r"\1"),
    (r"\\texttt\{([^}]*)\}", r"\1"),
    (r"\\url\{([^}]*)\}", r"\1"),
    (r"\\ref\{tab:taskspec\}", "1"),
    (r"\\ref\{tab:consortium\}", "2"),
    (r"\\ref\{fig:workflow\}", "1"),
    (r"\\ref\{fig:degradation\}", "2"),
    (r"\\ref\{fig:forest\}", "3"),
    (r"\\ref\{tab:supp-baseline-full\}", "S1"),
    (r"\\ref\{tab:supp-leaderboard\}", "S2"),
    (r"\\ref\{tab:mimic-snapshot\}", "S3"),
    (r"\\ref\{tab:comparators\}", "S4"),
    (r"\\ref\{([^}]*)\}", r"\1"),
    (r"\\label\{[^}]*\}", ""),
    (r"\\noindent", ""),
    (r"\\appendix", ""),
    (r"\\setcounter\{[^}]*\}\{[^}]*\}", ""),
    (r"\\renewcommand\{[^}]*\}\{[^}]*\}", ""),
    (r"\\bibliographystyle\{[^}]*\}", ""),
    (r"\\bibliography\{[^}]*\}", ""),
    (r"\\%", "%"),
    (r"\\_", "_"),
    (r"\\&", "&"),
    (r"\\#", "#"),
    (r"\\ge", ">="),
    (r"\\le", "<="),
    (r"\$([^$]*)\$", r"\1"),
    (r"\{", ""),
    (r"\}", ""),
]

BLOCK_RE = re.compile(r"\\begin\{(table|figure)\}.*?\\end\{\1\}", re.DOTALL)


def strip_latex(text: str) -> str:
    text = re.sub(r"\\[`'\"^~=.]\\?\{?([A-Za-z])\}?", r"\1", text)
    text = text.replace("\\&", "&").replace("\\_", "_").replace("\\%", "%")
    text = text.replace("{", "").replace("}", "")
    text = re.sub(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?", "", text)
    text = text.replace("~", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def citation_order() -> list[str]:
    order: list[str] = []
    text = SRC.read_text()
    for match in re.finditer(r"\\cite(?:p|t)?\{([^}]*)\}", text):
        for key in match.group(1).split(","):
            key = key.strip()
            if key and key not in order:
                order.append(key)
    return order


def parse_bib_entries() -> dict[str, dict[str, str]]:
    raw = BIB.read_text()
    entries: dict[str, dict[str, str]] = {}
    starts = list(re.finditer(r"@(\w+)\s*\{\s*([^,]+),", raw))
    for i, start in enumerate(starts):
        key = start.group(2).strip()
        end = starts[i + 1].start() if i + 1 < len(starts) else len(raw)
        block = raw[start.end() : end]
        fields: dict[str, str] = {"entrytype": start.group(1).lower()}
        for field in re.finditer(r"(\w+)\s*=\s*(\{([^{}]*(?:\{[^{}]*\}[^{}]*)*)\}|\"([^\"]*)\")", block, re.DOTALL):
            name = field.group(1).lower()
            val = field.group(3) if field.group(3) is not None else field.group(4)
            fields[name] = strip_latex(val or "")
        entries[key] = fields
    return entries


def short_authors(authors: str) -> str:
    if not authors:
        return ""
    parts = [a.strip() for a in authors.split(" and ") if a.strip()]
    if not parts:
        return authors
    if len(parts) > 3 or any(a.lower() == "others" for a in parts):
        first = parts[0]
        if "," in first:
            surname = first.split(",", 1)[0].strip()
        else:
            surname = first.split()[-1]
        return f"{surname} et al."
    rendered = []
    for author in parts:
        if "," in author:
            surname, given = [x.strip() for x in author.split(",", 1)]
            initials = "".join(p[0] for p in given.replace(".", " ").split() if p)
            rendered.append(f"{surname} {initials}")
        else:
            bits = author.split()
            rendered.append(bits[-1] + " " + "".join(b[0] for b in bits[:-1]))
    return ", ".join(rendered)


def format_reference(fields: dict[str, str], key: str) -> str:
    authors = short_authors(fields.get("author", ""))
    title = fields.get("title", key).rstrip(".")
    year = fields.get("year", "n.d.")
    venue = fields.get("journal") or fields.get("booktitle") or fields.get("howpublished") or fields.get("institution", "")
    volume = fields.get("volume", "")
    number = fields.get("number", "")
    pages = fields.get("pages", "")
    doi = fields.get("doi", "")
    note = fields.get("note", "")

    parts = []
    if authors:
        parts.append(authors + ".")
    parts.append(title + ".")
    if venue:
        venue_part = venue
        if volume:
            venue_part += f";{volume}"
            if number:
                venue_part += f"({number})"
        if pages:
            venue_part += f":{pages}"
        venue_part += f". {year}."
        parts.append(venue_part)
    else:
        parts.append(f"{year}.")
    if doi:
        parts.append(f"doi:{doi}.")
    elif note:
        parts.append(note.rstrip(".") + ".")
    return " ".join(parts)


def prepare_citations_and_references():
    global CITE_NUMBERS, REFERENCES
    order = citation_order()
    entries = parse_bib_entries()
    CITE_NUMBERS = {key: str(i) for i, key in enumerate(order, start=1)}
    REFERENCES = [format_reference(entries.get(key, {}), key) for key in order]


def citation_text(keys: str) -> str:
    numbers = [CITE_NUMBERS.get(key.strip(), key.strip()) for key in keys.split(",")]
    return "(" + ",".join(numbers) + ")"


def latex_to_text(text: str) -> str:
    text = re.sub(r"(?<!\\)%.*", "", text)
    text = text.replace("--", "-")
    text = re.sub(r"\\citep\{([^}]*)\}", lambda m: citation_text(m.group(1)), text)
    text = re.sub(r"\\cite\{([^}]*)\}", lambda m: citation_text(m.group(1)), text)
    text = re.sub(r"\\citet\{([^}]*)\}", lambda m: citation_text(m.group(1)), text)
    for pat, rep in SUBS:
        text = re.sub(pat, rep, text)
    text = re.sub(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?", "", text)
    text = text.replace("~", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def document_body() -> str:
    text = SRC.read_text()
    m = re.search(r"\\begin\{document\}(.*?)\\end\{document\}", text, re.DOTALL)
    if not m:
        raise ValueError("main_v3.tex has no document environment")
    body = m.group(1)
    body = re.sub(r"\\maketitle", "", body)
    return body


def set_run_typeface(run, size: float | None = None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.underline = False
    if size is not None:
        run.font.size = Pt(size)


def format_paragraph(p, size: float = 12, space_after: float = 6, double: bool = True):
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(space_after)
    if double:
        fmt.line_spacing = 2.0
    else:
        fmt.line_spacing = 1.0
    for run in p.runs:
        set_run_typeface(run, size)
    return p


def add_para(doc: Document, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead + " ")
        r.bold = True
    p.add_run(text)
    format_paragraph(p)
    return p


def add_table(doc: Document, spec: dict):
    cap = doc.add_paragraph()
    run = cap.add_run(spec["caption"])
    run.bold = True
    format_paragraph(cap, size=10, space_after=3, double=False)

    table = doc.add_table(rows=1, cols=len(spec["header"]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, heading in enumerate(spec["header"]):
        cell = table.rows[0].cells[i]
        cell.text = heading
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            set_run_typeface(run, 8)
        cell.paragraphs[0].paragraph_format.line_spacing = 1.0
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    for row in spec["rows"]:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cells[i].paragraphs[0].runs:
                set_run_typeface(run, 8)
            cells[i].paragraphs[0].paragraph_format.line_spacing = 1.0
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
    spacer = doc.add_paragraph()
    format_paragraph(spacer, size=12, space_after=0, double=False)


def add_figure(doc: Document, spec: dict):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    path = FIG / spec["image"]
    if path.exists():
        p.add_run().add_picture(str(path), width=Inches(spec["width"]))
    else:
        p.add_run(f"[missing figure: {spec['image']}]")
    format_paragraph(p, space_after=3, double=False)
    cap = doc.add_paragraph()
    run = cap.add_run(spec["caption"])
    run.italic = True
    format_paragraph(cap, size=10, space_after=3, double=False)
    spacer = doc.add_paragraph()
    format_paragraph(spacer, size=12, space_after=0, double=False)


def route_float(doc: Document, block: str):
    m = re.search(r"\\label\{([^}]*)\}", block)
    label = m.group(1) if m else ""
    if label in TABLES:
        add_table(doc, TABLES[label])
    elif label in FIGURES:
        add_figure(doc, FIGURES[label])


def emit_prose(doc: Document, block: str):
    block = re.sub(r"\\begin\{abstract\}", r"\\section*{Abstract}", block)
    block = re.sub(r"\\end\{abstract\}", "", block)
    block = re.sub(r"\\begin\{(itemize|enumerate)\}(\[[^\]]*\])?", "", block)
    block = re.sub(r"\\end\{(itemize|enumerate)\}", "", block)
    block = re.sub(r"\\item\s+", "\n\n- ", block)

    parts = re.split(r"(\\(?:section|subsection)\*?\{[^}]+\})", block)
    for part in parts:
        if not part.strip():
            continue
        heading = re.match(r"\\(section|subsection)\*?\{([^}]+)\}", part.strip())
        if heading:
            level = 1 if heading.group(1) == "section" else 2
            p = doc.add_paragraph()
            run = p.add_run(latex_to_text(heading.group(2)))
            run.bold = True
            if level == 2:
                run.italic = True
            format_paragraph(p, size=12, space_after=3, double=True)
            continue
        for para in re.split(r"\n\s*\n", part):
            para = para.strip()
            if not para:
                continue
            lead = re.match(r"\\textbf\{([^}]*)\}\s*(.*)", para, re.DOTALL)
            if lead:
                text = latex_to_text(lead.group(2))
                if text:
                    add_para(doc, text, bold_lead=latex_to_text(lead.group(1)))
            else:
                text = latex_to_text(para)
                if text:
                    add_para(doc, text)


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_typeface(run, 12)


def apply_base_format(doc: Document):
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    add_page_number(section)

    for style_name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
    doc.styles["Normal"].font.size = Pt(12)


def add_title_page(doc: Document):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(TITLE)
    run.bold = True
    format_paragraph(title, size=14, space_after=12, double=False)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(AUTHORS)
    format_paragraph(sub, size=12, space_after=12, double=False)

    for line in TITLE_PAGE_TODOS:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_paragraph(p, size=12, space_after=0, double=False)

    doc.add_page_break()


def add_references(doc: Document):
    p = doc.add_paragraph()
    r = p.add_run("References")
    r.bold = True
    format_paragraph(p, size=12, space_after=3, double=True)
    for idx, ref in enumerate(REFERENCES, start=1):
        p = doc.add_paragraph(f"{idx}. {ref}")
        format_paragraph(p, size=12, space_after=3, double=True)


def build_docx():
    prepare_citations_and_references()
    doc = Document()
    apply_base_format(doc)
    add_title_page(doc)

    body = document_body()
    pos = 0
    for match in BLOCK_RE.finditer(body):
        emit_prose(doc, body[pos:match.start()])
        route_float(doc, match.group(0))
        pos = match.end()
    emit_prose(doc, body[pos:])
    add_references(doc)

    doc.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build_docx()
